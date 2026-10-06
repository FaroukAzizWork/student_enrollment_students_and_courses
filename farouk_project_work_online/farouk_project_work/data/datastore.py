
import docx
import os


class WordDatabase:
    def __init__(self, file_path, headers):
        self.file_path = file_path
        self.headers = headers
        self._initialize_file()

    def _initialize_file(self):

        if not os.path.exists(self.file_path):
            doc = docx.Document()
            table = doc.add_table(rows=1, cols=len(self.headers))

            # تعبئة الصف الأول بأسماء الأعمدة
            hdr_cells = table.rows[0].cells
            for i in range(len(self.headers)):
                hdr_cells[i].text = self.headers[i]

            doc.save(self.file_path)

    def add_row(self, row_data):
        """تضيف صفاً جديداً إلى الجدول وتحفظ الملف"""
        doc = docx.Document(self.file_path)
        table = doc.tables[0]  # نتعامل دائماً مع الجدول الأول في الملف

        row_cells = table.add_row().cells
        for i in range(len(row_data)):
            row_cells[i].text = str(row_data[i])

        doc.save(self.file_path)

    def get_all_rows(self):
        """تقرأ كل البيانات من الجدول وترجعها كقائمة"""
        if not os.path.exists(self.file_path):
            return []

        doc = docx.Document(self.file_path)
        table = doc.tables[0]
        data = []

        # نتجاهل الصف الأول لأنه يحتوي على أسماء الأعمدة فقط
        for row in table.rows[1:]:
            row_data = [cell.text for cell in row.cells]
            data.append(row_data)

        return data